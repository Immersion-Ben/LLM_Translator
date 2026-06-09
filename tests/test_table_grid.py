from table_grid import parse_table_html, TableGrid, GridCell


def _cell(g, row, col):
    for c in g.cells:
        if c.row == row and c.col == col:
            return c
    raise AssertionError(f"no origin cell at ({row},{col})")


def test_simple_grid():
    html = "<table><tr><td>A</td><td>B</td></tr><tr><td>C</td><td>D</td></tr></table>"
    g = parse_table_html(html)
    assert (g.n_rows, g.n_cols) == (2, 2)
    assert _cell(g, 1, 1).text == "D"


def test_colspan_shifts_following_cell():
    html = "<table><tr><td colspan='3'>T</td><td>X</td></tr></table>"
    g = parse_table_html(html)
    assert g.n_cols == 4
    assert _cell(g, 0, 0).colspan == 3
    assert _cell(g, 0, 3).text == "X"


def test_rowspan_tracks_column():
    html = "<table><tr><td rowspan='2'>R</td><td>a</td></tr><tr><td>b</td></tr></table>"
    g = parse_table_html(html)
    assert (g.n_rows, g.n_cols) == (2, 2)
    assert _cell(g, 1, 1).text == "b"


def test_tolerates_unclosed_tbody():
    html = "<html><body><table><tr><td>A</td><td>B</td></tr></tbody></table></body></html>"
    g = parse_table_html(html)
    assert (g.n_rows, g.n_cols) == (1, 2)


from docx import Document
from table_grid import build_docx_table


def test_build_docx_merges_and_translates():
    g = parse_table_html("<table><tr><td colspan='2'>Header</td></tr><tr><td>a</td><td>b</td></tr></table>")
    doc = Document()
    table = build_docx_table(doc, g, translate=lambda s: s.upper())
    assert (len(table.rows), len(table.columns)) == (2, 2)
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc  # 병합됨
    assert table.cell(0, 0).text == "HEADER"
    assert table.cell(1, 1).text == "B"
